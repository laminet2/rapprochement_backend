import uuid
import unicodedata
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from django.conf import settings
from django.db import transaction
from django.utils import timezone

from rest_framework import generics, status
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.response import Response

from rest_framework.views import APIView

from .models import TypeTransactionCaisse, TransactionValorix

from .serializers import (
    ImportValorixSerializer,
    TransactionValorixSerializer,
    TransactionValorixJSONSerializer
)
from core.models import Rapprochement,StatutValidation,TransactionValorixRapprochee

from rest_framework.permissions import IsAuthenticated


from core.services.actions import (
    annuler_rapprochement,
    ActionRapprochementError,
)

class ImportValorixView(generics.GenericAPIView):
    serializer_class = ImportValorixSerializer
    parser_classes = [MultiPartParser, FormParser]

    MAX_ERRORS = 100
    BATCH_SIZE = 500

    # Indices des colonnes Excel (indexation Python à partir de zéro).
    COL_TRANSACTION = 1   # B : N° Trn
    COL_DATE = 4          # E : Date et Heure
    COL_COMPTE = 8        # I : N° Compte
    COL_NOM = 9           # J : Nom & Prénom
    COL_CREDIT = 12       # M : Crédit
    COL_DEBIT = 13        # N : Débit
    COL_REFERENCES = 14   # O : Références

    # Les catégories du fichier peuvent comporter des accents
    # ou des différences de casse.
    CATEGORIES = {
        "depot cheques": TypeTransactionCaisse.DEPOT_CHEQUE,
        "retrait cheques": TypeTransactionCaisse.RETRAIT_CHEQUE,
        "others": TypeTransactionCaisse.AUTRE,
        "depot especes": TypeTransactionCaisse.DEPOT_ESPECES,
        "retrait especes": TypeTransactionCaisse.RETRAIT_ESPECES,
        "retrait especes carte visa cgf":
            TypeTransactionCaisse.RETRAIT_ESPECES,
        "transferts": TypeTransactionCaisse.TRANSFERT,
        "depot virements": TypeTransactionCaisse.DEPOT_VIREMENTS,
        "retrait virements": TypeTransactionCaisse.RETRAIT_VIREMENTS,
    }

    @staticmethod
    def normaliser_texte(valeur):
        """Supprime les accents et normalise les espaces et la casse."""
        texte = str(valeur or "").strip().lower()
        texte = unicodedata.normalize("NFKD", texte)
        texte = "".join(
            caractere
            for caractere in texte
            if not unicodedata.combining(caractere)
        )
        return " ".join(texte.split())

    @staticmethod
    def convertir_montant(valeur):
        """Convertit un montant Excel en Decimal."""

        if valeur is None:
            return None

        if isinstance(valeur, str):
            texte = valeur.strip()

            if not texte:
                return None

            texte = texte.replace("\xa0", "").replace(" ", "")

            # Formats français : 1234,56 ou 1.234,56.
            if "," in texte and "." in texte:
                texte = texte.replace(".", "").replace(",", ".")
            elif "," in texte:
                texte = texte.replace(",", ".")

            valeur = texte

        try:
            montant = Decimal(str(valeur))
        except (InvalidOperation, ValueError, TypeError):
            raise ValueError(
                f"Montant invalide : {valeur!r}"
            )

        if not montant.is_finite():
            raise ValueError(
                f"Le montant doit être un nombre fini : {valeur!r}"
            )

        return montant

    @staticmethod
    def convertir_date(valeur, epoch):
        """Convertit une date Excel ou une chaîne en datetime."""

        if valeur is None or valeur == "":
            raise ValueError("Date de transaction manquante.")

        if isinstance(valeur, datetime):
            resultat = valeur

        elif isinstance(valeur, date):
            resultat = datetime.combine(valeur, time.min)

        elif isinstance(valeur, (int, float)):
            # Excel peut représenter une date sous forme de nombre sériel.
            resultat = from_excel(valeur, epoch=epoch)

            if isinstance(resultat, time):
                raise ValueError(
                    "La cellule contient une heure sans date."
                )

            if isinstance(resultat, date) and not isinstance(
                resultat, datetime
            ):
                resultat = datetime.combine(resultat, time.min)

            if not isinstance(resultat, datetime):
                raise ValueError("Date Excel invalide.")

        elif isinstance(valeur, str):
            texte = valeur.strip()
            resultat = None

            formats = (
                "%d/%m/%Y %H:%M:%S.%f",
                "%d/%m/%Y %H:%M:%S",
                "%d/%m/%Y %H:%M",
                "%d/%m/%Y",
                "%Y-%m-%d %H:%M:%S.%f",
                "%Y-%m-%d %H:%M:%S",
                "%Y-%m-%d %H:%M",
                "%Y-%m-%d",
                "%d-%m-%Y %H:%M:%S",
                "%d-%m-%Y",
            )

            for fmt in formats:
                try:
                    resultat = datetime.strptime(texte, fmt)
                    break
                except ValueError:
                    continue

            if resultat is None:
                raise ValueError(
                    f"Format de date invalide : {texte!r}"
                )

        else:
            raise ValueError(
                f"Type de date non pris en charge : "
                f"{type(valeur).__name__}"
            )

        # Ajoute le fuseau horaire uniquement si Django utilise les fuseaux.
        if settings.USE_TZ and timezone.is_naive(resultat):
            resultat = timezone.make_aware(
                resultat,
                timezone.get_current_timezone(),
            )

        return resultat

    @staticmethod
    def convertir_save(valeur):
        """Convertit explicitement le paramètre save en booléen."""

        if isinstance(valeur, bool):
            return valeur

        texte = str(valeur if valeur is not None else "false")
        texte = texte.strip().lower()

        if texte in ("true", "1", "yes", "oui"):
            return True

        if texte in ("false", "0", "no", "non", ""):
            return False

        raise ValueError(
            "Le paramètre save doit être true ou false."
        )

    def post(self, request, *args, **kwargs):
        # 1. Validation du fichier envoyé.
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)

        fichier = serializer.validated_data["fichier"]

        try:
            save = self.convertir_save(request.data.get("save", False))
        except ValueError as erreur:
            return Response(
                {"detail": str(erreur)},
                status=status.HTTP_400_BAD_REQUEST,
            )

        # 2. Ouverture du classeur Excel.
        try:
            fichier.seek(0)

            workbook = load_workbook(
                fichier,
                read_only=True,
                data_only=True,
            )
        except Exception:
            return Response(
                {"detail": "Le fichier Excel est invalide ou illisible."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        try:
            feuille = workbook.active

            # Dans le fichier fourni :
            # ligne 12 = en-têtes
            # ligne 13 = première catégorie
            # ligne 14 = premières transactions
            lignes = feuille.iter_rows(
                min_row=13,
                values_only=True,
            )

            transactions = []
            erreurs = []
            type_courant = None
            erreurs_tronquees = False

            # 3. Lecture des lignes Excel.
            for numero_ligne, ligne in enumerate(lignes, start=13):

                # Ignore les lignes entièrement vides.
                if not any(valeur is not None for valeur in ligne):
                    continue

                # Vérifie que la ligne contient les colonnes requises.
                derniere_colonne = max(
                    self.COL_TRANSACTION,
                    self.COL_DATE,
                    self.COL_COMPTE,
                    self.COL_NOM,
                    self.COL_CREDIT,
                    self.COL_DEBIT,
                    self.COL_REFERENCES,
                )

                if len(ligne) <= derniere_colonne:
                    if len(erreurs) < self.MAX_ERRORS:
                        erreurs.append({
                            "ligne": numero_ligne,
                            "erreur": (
                                "La ligne ne contient pas toutes "
                                "les colonnes attendues."
                            ),
                        })
                    else:
                        erreurs_tronquees = True
                        break

                    continue

                numero = ligne[self.COL_TRANSACTION]

                # Détection des catégories en colonne B.
                if isinstance(numero, str):
                    categorie = self.normaliser_texte(numero)

                    if categorie in self.CATEGORIES:
                        type_courant = self.CATEGORIES[categorie]

                    # Ignore également les lignes "Nb Opérations",
                    # "Nombre", les en-têtes et autres libellés.
                    continue

                # Les lignes de synthèse ne sont pas des transactions.
                if not isinstance(numero, (int, float)):
                    continue

                # Les numéros de transaction doivent être entiers.
                if isinstance(numero, float) and not numero.is_integer():
                    erreur = ValueError(
                        "Le numéro de transaction doit être entier."
                    )

                    if len(erreurs) < self.MAX_ERRORS:
                        erreurs.append({
                            "ligne": numero_ligne,
                            "erreur": str(erreur),
                        })
                    else:
                        erreurs_tronquees = True
                        break

                    continue

                try:
                    if type_courant is None:
                        raise ValueError(
                            "Type de transaction non reconnu."
                        )

                    date_transaction = self.convertir_date(
                        ligne[self.COL_DATE],
                        workbook.epoch,
                    )

                    debit = self.convertir_montant(
                        ligne[self.COL_DEBIT]
                    )

                    credit = self.convertir_montant(
                        ligne[self.COL_CREDIT]
                    )

                    numero_compte = ligne[self.COL_COMPTE]
                    nom = ligne[self.COL_NOM]
                    references = ligne[self.COL_REFERENCES]

                    # Création en mémoire ; aucune écriture en base ici.
                    transactions.append(
                        TransactionValorix(
                            id=uuid.uuid4(),
                            numero_transaction=str(int(numero)),
                            date_transaction=date_transaction,
                            numero_compte=str(
                                numero_compte or ""
                            ).strip(),
                            nomComplet=str(nom or "").strip(),
                            debit=debit,
                            credit=credit,
                            references=str(
                                references or ""
                            ).strip(),
                            types=type_courant,
                        )
                    )

                except (
                    ValueError,
                    TypeError,
                    InvalidOperation,
                    IndexError,
                    OverflowError,
                ) as erreur:
                    if len(erreurs) < self.MAX_ERRORS:
                        erreurs.append({
                            "ligne": numero_ligne,
                            "erreur": str(erreur),
                        })
                    else:
                        erreurs_tronquees = True
                        break

            # 4. Si une erreur existe, on n'enregistre aucune transaction.
            if erreurs:
                return Response(
                    {
                        "detail": (
                            "Importation annulée : des erreurs "
                            "ont été détectées."
                        ),
                        "nombre_erreurs_retournees": len(erreurs),
                        "erreurs": erreurs,
                        "erreurs_tronquees": erreurs_tronquees,
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            # 5. Vérifie que le fichier contient des transactions.
            if not transactions:
                return Response(
                    {
                        "detail": (
                            "Aucune transaction trouvée dans le fichier."
                        ),
                        "nombre_transactions": 0,
                        "enregistre": False,
                        "data": [],
                    },
                    status=status.HTTP_200_OK,
                )

            # 6. Enregistrement facultatif et atomique.
            if save:
                with transaction.atomic():
                    TransactionValorix.objects.bulk_create(
                        transactions,
                        batch_size=self.BATCH_SIZE,
                    )

            # 7. Sérialisation des données pour la réponse API.
            resultat = TransactionValorixSerializer(
                transactions,
                many=True,
            )

            return Response(
                {
                    "detail": (
                        "Importation enregistrée avec succès."
                        if save
                        else (
                            "Lecture réussie. "
                            "Aucune donnée n'a été enregistrée."
                        )
                    ),
                    "nombre_transactions": len(transactions),
                    "enregistre": save,
                    "data": resultat.data,
                },
                status=(
                    status.HTTP_201_CREATED
                    if save
                    else status.HTTP_200_OK
                ),
            )

        finally:
            workbook.close()



class ImportValorixJSONView(APIView):
    
    def post(self, request, *args, **kwargs):
        # Accepte soit une liste JSON, soit un objet contenant
        # une clé "transactions".
        donnees = request.data

        if isinstance(donnees, dict):
            donnees = donnees.get("transactions")

        if not isinstance(donnees, list):
            return Response(
                {
                    "detail": (
                        "Le JSON doit être une liste de transactions "
                        "ou un objet contenant la clé 'transactions'."
                    )
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        if not donnees:
            return Response(
                {"detail": "Aucune transaction à enregistrer."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        serializer = TransactionValorixJSONSerializer(
            data=donnees,
            many=True,
        )

        serializer.is_valid(raise_exception=True)

        transactions = []

        for item in serializer.validated_data:
            transaction_data = dict(item)

            # Si l'identifiant est fourni, on le conserve.
            # Sinon, on génère un UUID.
            transaction_id = transaction_data.pop(
                "id",
                uuid.uuid4(),
            )

            transactions.append(
                TransactionValorix(
                    id=transaction_id,
                    **transaction_data,
                )
            )

        try:
            with transaction.atomic():
                TransactionValorix.objects.bulk_create(
                    transactions,
                    batch_size=500,
                )

        except Exception as erreur:
            # Les détails techniques sont conservés dans les logs,
            # mais pas exposés au client.
            import logging

            logging.getLogger(__name__).exception(
                "Erreur lors de l'importation JSON Valorix."
            )

            return Response(
                {
                    "detail": (
                        "Impossible d'enregistrer les transactions. "
                        "Vérifiez les identifiants et les données envoyées."
                    )
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        return Response(
            {"message": "Transactions bien enregistrées."},
            status=status.HTTP_201_CREATED,
        )

class TransactionListView(generics.ListAPIView):
    queryset = TransactionValorix.objects.all().order_by(
        "-date_transaction"
    )
    serializer_class = TransactionValorixSerializer
