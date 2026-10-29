"""
Lecture des documents de contrepartie (Excel, CSV ou PDF) et enregistrement
dans un lot d'import. Les colonnes obligatoires sont : numéro compte client,
date transaction, montant net.
"""

import csv
import io
from collections import Counter
from datetime import datetime
from decimal import Decimal

from django.db import transaction as db_transaction

from documents.models import LigneDocument, SensDocument
from documents.prompts import PROMPT_DOCUMENT
from documents.schemas import ExtractionDocument
from importation.models import LotImport, SourceImport
from importation.outils import (
    a_zero,
    convertir_date,
    convertir_montant,
    marquer_doublons,
    montant_positif,
    montant_vers_texte,
    normaliser_entete,
)
from importation.pdf import (
    ExtractionError,
    extraire_par_blocs,
    lire_pages_pdf,
    regrouper_pages,
)

COLONNES_REQUISES = {
    "compte": "Numéro compte client",
    "date": "Date transaction",
    "montant": "Montant net",
}

_EXACT = {
    "compte": {
        "numero compte client", "numero de compte client", "n compte client",
        "no compte client", "compte client", "numero compte",
        "numero de compte", "n compte", "num compte", "compte",
    },
    "date": {"date transaction", "date de transaction", "date operation", "date"},
    "montant": {"montant net", "net"},
}


def _texte(valeur):
    if isinstance(valeur, float) and valeur.is_integer():
        valeur = int(valeur)
    return str(valeur if valeur is not None else "").strip()


def detecter_colonnes(entetes):
    """Associe les en-têtes aux 3 colonnes requises -> (colonnes, manquantes)."""
    normalises = [normaliser_entete(e) for e in entetes]
    colonnes, utilises = {}, set()

    def attribuer(cle, predicats):
        for predicat in predicats:
            for i, texte in enumerate(normalises):
                if i not in utilises and texte and predicat(texte):
                    colonnes[cle] = i
                    utilises.add(i)
                    return

    attribuer("montant", [
        lambda t: t in _EXACT["montant"],
        lambda t: "montant" in t and "net" in t,
        lambda t: "montant" in t,
    ])
    attribuer("compte", [lambda t: t in _EXACT["compte"], lambda t: "compte" in t])
    attribuer("date", [
        lambda t: t in _EXACT["date"],
        lambda t: t.startswith("date"),
        lambda t: "date" in t,
    ])

    manquantes = [l for c, l in COLONNES_REQUISES.items() if c not in colonnes]
    return colonnes, manquantes


def construire_ligne(index, compte, date_val, montant, epoch=None):
    erreurs, alertes = [], []

    compte_txt = _texte(compte)
    if not compte_txt:
        erreurs.append("Numéro de compte client manquant.")

    jour = None
    try:
        jour = convertir_date(date_val, epoch).date().isoformat()
    except ValueError as exc:
        erreurs.append(str(exc))

    valeur = None
    try:
        valeur = convertir_montant(montant)
        if valeur is None:
            erreurs.append("Montant net manquant.")
        elif valeur < 0:
            erreurs.append("Montant net négatif non accepté.")
        elif valeur == 0:
            erreurs.append("Montant net nul.")
    except ValueError as exc:
        erreurs.append(str(exc))

    return {
        "index": index,
        "numero_compte": compte_txt,
        "date_transaction": jour,
        "montant": montant_vers_texte(valeur) if valeur is not None else None,
        "doublon_probable": False,
        "erreurs": erreurs,
        "alertes": alertes,
    }


def _lignes_depuis_tableau(brutes, epoch=None):
    colonnes, debut = None, None
    for i, ligne in enumerate(brutes[:30]):
        cols, manquantes = detecter_colonnes(ligne)
        if not manquantes:
            colonnes, debut = cols, i + 1
            break

    if colonnes is None:
        premiere = next((l for l in brutes[:30] if any(l)), [])
        _, manquantes = detecter_colonnes(premiere)
        trouvees = ", ".join(str(e) for e in premiere if e) or "aucune"
        raise ExtractionError(
            f"Colonnes manquantes : {', '.join(manquantes)}. "
            f"Colonnes détectées dans le fichier : {trouvees}."
        )

    lignes = []
    for ligne in brutes[debut:]:
        def cellule(cle):
            i = colonnes[cle]
            return ligne[i] if i < len(ligne) else None

        compte, date_val, montant = cellule("compte"), cellule("date"), cellule("montant")
        if all(v in (None, "") for v in (compte, date_val, montant)):
            continue
        if date_val in (None, "") and (
            _texte(compte) == ""
            or normaliser_entete(compte).startswith(("total", "sous total"))
        ):
            continue  # ligne de total / commentaire
        lignes.append(construire_ligne(len(lignes) + 1, compte, date_val, montant, epoch))
    return lignes


def lire_document(fichier):
    """Fichier (.xlsx, .csv, .pdf) -> lignes à vérifier. N'écrit rien."""
    nom = fichier.name.lower()
    fichier.seek(0)

    if nom.endswith(".xlsx"):
        from openpyxl import load_workbook

        try:
            classeur = load_workbook(fichier, read_only=True, data_only=True)
        except Exception as exc:
            raise ExtractionError("Le fichier Excel est invalide ou illisible.") from exc
        try:
            brutes = [list(l) for l in classeur.active.iter_rows(values_only=True)]
            return _lignes_depuis_tableau(brutes, classeur.epoch)
        finally:
            classeur.close()

    if nom.endswith(".csv"):
        contenu = fichier.read()
        try:
            texte = contenu.decode("utf-8-sig")
        except UnicodeDecodeError:
            texte = contenu.decode("latin-1")
        try:
            dialecte = csv.Sniffer().sniff(texte[:4096], delimiters=",;\t|")
        except csv.Error:
            dialecte = csv.excel
            dialecte.delimiter = ";"
        brutes = [list(l) for l in csv.reader(io.StringIO(texte), dialecte)]
        return _lignes_depuis_tableau(brutes)

    if nom.endswith(".pdf"):
        pages = lire_pages_pdf(fichier)
        blocs = regrouper_pages(pages, pages_par_bloc=3)
        brutes = extraire_par_blocs(
            blocs, ExtractionDocument, PROMPT_DOCUMENT, "document de transactions",
            max_workers=3,
        )
        return [
            construire_ligne(i, b.numero_compte_client, b.date_transaction, b.montant_net)
            for i, b in enumerate(brutes, 1)
        ]

    raise ExtractionError("Format non pris en charge : utilisez .xlsx, .csv ou .pdf.")


def generer_modele():
    """Modèle Excel à télécharger : en-têtes + une ligne d'exemple."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.append(list(COLONNES_REQUISES.values()))
    ws.append(["0012345678", datetime.now().strftime("%d/%m/%Y"), 150000])
    for cellule in ws[1]:
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = PatternFill("solid", fgColor="1F3A5F")
    for lettre, largeur in zip("ABC", (26, 20, 18)):
        ws.column_dimensions[lettre].width = largeur
    ws["E1"] = "Supprimez la ligne d'exemple avant l'import."
    sortie = io.BytesIO()
    wb.save(sortie)
    return sortie.getvalue()


# ------------------------------------------------------------------
# Doublons
# ------------------------------------------------------------------

def cle_ligne(compte, jour, montant):
    return ((compte or "").strip(), jour, a_zero(montant))


def _compteur_existant(type_document, dates):
    if not dates:
        return Counter()
    qs = LigneDocument.objects.filter(
        type_document=type_document, date_transaction__range=(min(dates), max(dates))
    ).values_list("numero_compte", "date_transaction", "debit", "credit")
    return Counter(cle_ligne(c, d, deb or cred) for c, d, deb, cred in qs)


def marquer_doublons_lignes(lignes, type_document):
    valides = [l for l in lignes if l["date_transaction"] and l["montant"]]
    dates = [datetime.strptime(l["date_transaction"], "%Y-%m-%d").date() for l in valides]
    compteur = _compteur_existant(type_document, dates)
    cles = [
        cle_ligne(l["numero_compte"], d, Decimal(l["montant"]))
        for l, d in zip(valides, dates)
    ]
    for ligne, doublon in zip(valides, marquer_doublons(cles, compteur)):
        ligne["doublon_probable"] = doublon
    return lignes


# ------------------------------------------------------------------
# Enregistrement
# ------------------------------------------------------------------

@db_transaction.atomic
def enregistrer_import(lignes, type_document, nom_fichier, ignorer_doublons):
    """Enregistre des lignes validées. Retourne (lot, créées, ignorées)."""
    doublons = [False] * len(lignes)
    if ignorer_doublons:
        compteur = _compteur_existant(type_document, [l["date_transaction"] for l in lignes])
        doublons = marquer_doublons(
            [cle_ligne(l["numero_compte"], l["date_transaction"], l["montant"]) for l in lignes],
            compteur,
        )

    lot = LotImport.objects.create(
        source=SourceImport.DOCUMENT,
        nom_fichier=nom_fichier or "",
        details={"type_document": type_document.id, "nom": type_document.nom},
    )

    credit = type_document.sens == SensDocument.CREDIT
    a_creer, ignorees = [], 0
    for ligne, doublon in zip(lignes, doublons):
        if doublon:
            ignorees += 1
            continue
        montant = montant_positif(ligne["montant"])
        a_creer.append(
            LigneDocument(
                type_document=type_document,
                lot=lot,
                date_transaction=ligne["date_transaction"],
                numero_compte=ligne["numero_compte"].strip(),
                debit=None if credit else montant,
                credit=montant if credit else None,
            )
        )
    LigneDocument.objects.bulk_create(a_creer, batch_size=500)

    lot.nombre_lignes = len(a_creer)
    lot.nombre_doublons_ignores = ignorees
    lot.save(update_fields=["nombre_lignes", "nombre_doublons_ignores"])
    return lot, len(a_creer), ignorees
