import io
from datetime import date, datetime
from decimal import Decimal
from types import SimpleNamespace
from unittest import mock

from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook
from rest_framework.test import APITestCase

from core.tests.factories import creer_compte, tx_banque, tx_valorix
from importation.models import LotImport
from importation.outils import marquer_doublons
from releverBancaire.models import TransactionBancaire
from documents.importers import detecter_colonnes
from valorix.models import TransactionValorix


def xlsx_bytes(lignes):
    wb = Workbook()
    ws = wb.active
    for ligne in lignes:
        ws.append(ligne)
    sortie = io.BytesIO()
    wb.save(sortie)
    return sortie.getvalue()


def fichier(nom, contenu):
    return SimpleUploadedFile(nom, contenu)


def releve_valorix_xlsx():
    wb = Workbook()
    ws = wb.active

    def ecrire(row, **cols):
        for lettre, valeur in cols.items():
            ws[f"{lettre}{row}"] = valeur

    ecrire(12, B="N° Trn", E="Date et Heure")
    ecrire(13, B="Dépôt chèques")
    ecrire(14, B=1001, E=datetime(2026, 10, 5, 9, 30), I="0012345678", J="DIOP Awa",
           M=None, N=150000, O="CHQ 1234567")
    ecrire(15, B=1002, E="06/10/2026 11:00:00", I="0099887766", J="FALL Moussa",
           M=None, N="75 000,50", O="")
    ecrire(16, B="Nb Opérations")
    ecrire(17, B="Retrait espèces")
    ecrire(18, B=1003, E=datetime(2026, 10, 7), I="0012345678", J="DIOP Awa",
           M=20000, N=None, O="")
    ecrire(19, B=1004, E="pas une date", I="1", J="X", M=None, N=10, O="")
    sortie = io.BytesIO()
    wb.save(sortie)
    return sortie.getvalue()


class OutilsTests(APITestCase):
    def test_doublons_comptent_les_occurrences(self):
        flags = marquer_doublons(["a", "a", "a", "b"], {"a": 2})
        self.assertEqual(flags, [True, True, False, False])

    def test_detection_de_colonnes_tolerante(self):
        cols, manquantes = detecter_colonnes(
            ["N° de compte client", "Date de transaction", "Montant Net (XOF)"]
        )
        self.assertEqual(manquantes, [])
        self.assertEqual(cols, {"montant": 2, "compte": 0, "date": 1})

    def test_colonnes_manquantes(self):
        _, manquantes = detecter_colonnes(["Compte", "Libellé"])
        self.assertEqual(manquantes, ["Date transaction", "Montant net"])


class ImportValorixTests(APITestCase):
    def test_apercu_excel_ne_sauvegarde_rien(self):
        r = self.client.post(
            "/valorix/",
            {"fichier": fichier("releve.xlsx", releve_valorix_xlsx())},
            format="multipart",
        )
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(r.data["nombre_lignes"], 4)
        self.assertEqual(r.data["nombre_erreurs"], 1)  # la date invalide
        types = [l["types"] for l in r.data["data"]]
        self.assertEqual(types, ["DEPOT_CHEQUE", "DEPOT_CHEQUE", "RETRAIT_ESPECES", "RETRAIT_ESPECES"])
        self.assertEqual(r.data["data"][1]["debit"], "75000.50")
        self.assertEqual(TransactionValorix.objects.count(), 0)

    def test_format_non_supporte(self):
        r = self.client.post("/valorix/", {"fichier": fichier("a.txt", b"x")}, format="multipart")
        self.assertEqual(r.status_code, 400)

    def test_enregistrement_cree_un_lot_et_detecte_les_doublons(self):
        apercu = self.client.post(
            "/valorix/", {"fichier": fichier("r.xlsx", releve_valorix_xlsx())}, format="multipart"
        ).data
        valides = [l for l in apercu["data"] if not l["erreurs"]]

        r = self.client.post(
            "/valorix/save/",
            {"nom_fichier": "r.xlsx", "transactions": valides},
            format="json",
        )
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual(r.data["nombre_enregistrees"], 3)
        lot = LotImport.objects.get(pk=r.data["lot"])
        self.assertEqual(lot.source, "VALORIX")
        self.assertEqual(TransactionValorix.objects.filter(lot=lot).count(), 3)

        # Même fichier réimporté : tout est signalé comme doublon probable
        apercu2 = self.client.post(
            "/valorix/", {"fichier": fichier("r.xlsx", releve_valorix_xlsx())}, format="multipart"
        ).data
        self.assertEqual(apercu2["nombre_doublons"], 3)

        r2 = self.client.post(
            "/valorix/save/",
            {"ignorer_doublons": True, "transactions": valides},
            format="json",
        )
        self.assertEqual(r2.data["nombre_enregistrees"], 0)
        self.assertEqual(r2.data["nombre_ignorees_doublons"], 3)
        self.assertEqual(TransactionValorix.objects.count(), 3)

    def test_le_champ_obsolete_statut_rapprochement_est_ignore(self):
        r = self.client.post(
            "/valorix/save/",
            {"transactions": [{
                "numero_transaction": "9", "date_transaction": "2026-10-05T10:00:00Z",
                "debit": "100", "types": "DEPOT_CHEQUE", "statut_rapprochement": False,
            }]},
            format="json",
        )
        self.assertEqual(r.status_code, 201, r.content)

    def test_debit_et_credit_simultanes_refuses(self):
        r = self.client.post(
            "/valorix/save/",
            {"transactions": [{
                "numero_transaction": "9", "date_transaction": "2026-10-05T10:00:00Z",
                "debit": "100", "credit": "50", "types": "DEPOT_CHEQUE",
            }]},
            format="json",
        )
        self.assertEqual(r.status_code, 400)


class ImportBanqueTests(APITestCase):
    def setUp(self):
        self.compte = creer_compte()

    def faux_extraction(self):
        from releverBancaire.schemas import LigneTransactionSchema

        lignes = [
            LigneTransactionSchema(date_transaction="2026-10-05", libelle_brut="REMISE CHQ 1234567",
                                   credit=150000, numero_compte_ou_cheque="1234567"),
            LigneTransactionSchema(date_transaction="2026-10-06", libelle_brut="FRAIS", debit=500),
            LigneTransactionSchema(date_transaction="5 octobre", libelle_brut="DATE ILLISIBLE", debit=1),
            LigneTransactionSchema(date_transaction="2026-10-06", libelle_brut="DEUX MONTANTS",
                                   debit=5, credit=5),
        ]
        return mock.patch.multiple(
            "releverBancaire.importers",
            lire_pages_pdf=mock.Mock(return_value=[(1, "texte")]),
            extraire_par_blocs=mock.Mock(return_value=lignes),
        )

    def test_apercu_pdf_signale_les_lignes_douteuses_sans_rien_enregistrer(self):
        with self.faux_extraction():
            r = self.client.post(
                f"/rbc/compte/{self.compte.id}/",
                {"fichier": fichier("releve.pdf", b"%PDF")}, format="multipart",
            )
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(r.data["nombre_lignes"], 4)
        self.assertEqual(r.data["nombre_erreurs"], 2)
        self.assertEqual(r.data["data"][0]["credit"], "150000.00")
        self.assertEqual(TransactionBancaire.objects.count(), 0)

    def test_compte_inconnu(self):
        r = self.client.post("/rbc/compte/999/", {"fichier": fichier("a.pdf", b"x")}, format="multipart")
        self.assertEqual(r.status_code, 404)

    def test_enregistrement_lot_et_doublons(self):
        corps = {
            "compte": self.compte.id, "nom_fichier": "releve.pdf", "ignorer_doublons": True,
            "transactions": [
                {"date_transaction": "2026-10-05", "libelle_brut": "REMISE", "credit": "1000"},
                {"date_transaction": "2026-10-05", "libelle_brut": "REMISE", "credit": "1000"},
                {"date_transaction": "2026-10-06", "libelle_brut": "FRAIS", "debit": "50"},
            ],
        }
        r = self.client.post("/rbc/save/", corps, format="json")
        self.assertEqual(r.status_code, 201, r.content)
        # Deux lignes identiques DANS le fichier ne sont pas des doublons.
        self.assertEqual(r.data["nombre_enregistrees"], 3)

        # Réimport du même relevé : tout est doublon.
        r2 = self.client.post("/rbc/save/", corps, format="json")
        self.assertEqual(r2.data["nombre_enregistrees"], 0)
        self.assertEqual(r2.data["nombre_ignorees_doublons"], 3)
        self.assertEqual(TransactionBancaire.objects.count(), 3)

    def test_apercu_marque_les_doublons(self):
        tx_banque(date(2026, 10, 5), 150000, "CREDIT", libelle="REMISE CHQ 1234567", compte=self.compte)
        with self.faux_extraction():
            r = self.client.post(
                f"/rbc/compte/{self.compte.id}/",
                {"fichier": fichier("releve.pdf", b"%PDF")}, format="multipart",
            )
        self.assertEqual(r.data["nombre_doublons"], 1)
        self.assertTrue(r.data["data"][0]["doublon_probable"])

    def test_compte_obligatoire(self):
        r = self.client.post(
            "/rbc/save/",
            {"transactions": [{"date_transaction": "2026-10-05", "libelle_brut": "X", "credit": "1"}]},
            format="json",
        )
        self.assertEqual(r.status_code, 400)

    def test_ancien_format_liste_avec_compte_par_ligne(self):
        r = self.client.post(
            "/rbc/save/",
            [{"compte": self.compte.id, "date_transaction": "2026-10-05",
              "libelle_brut": "X", "credit": "1"}],
            format="json",
        )
        self.assertEqual(r.status_code, 201, r.content)

    def test_comptes_et_banques(self):
        r = self.client.get("/rbc/comptes/")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.data[0]["numero_compte"], "COMPTE-001")


class LotsTests(APITestCase):
    def test_annulation_d_un_lot(self):
        r = self.client.post(
            "/valorix/save/",
            {"transactions": [{
                "numero_transaction": "1", "date_transaction": "2026-10-05T10:00:00Z",
                "debit": "100", "types": "DEPOT_CHEQUE"}]},
            format="json",
        )
        lot_id = r.data["lot"]
        liste = self.client.get("/imports/")
        self.assertEqual(liste.data["count"], 1)

        a = self.client.post(f"/imports/{lot_id}/annuler/")
        self.assertEqual(a.status_code, 200, a.content)
        self.assertEqual(TransactionValorix.objects.count(), 0)
        self.assertIsNotNone(LotImport.objects.get(pk=lot_id).annule_le)
        self.assertEqual(self.client.post(f"/imports/{lot_id}/annuler/").status_code, 409)

    def test_annulation_refusee_si_une_ligne_est_rapprochee(self):
        from core.services.manuel import creer_rapprochement_manuel
        lot = LotImport.objects.create(source="VALORIX")
        v = tx_valorix(date(2026, 10, 5), 100, "DEBIT", lot=lot)
        b = tx_banque(date(2026, 10, 5), 100, "CREDIT")
        creer_rapprochement_manuel([b.pk], [v.pk])
        a = self.client.post(f"/imports/{lot.id}/annuler/")
        self.assertEqual(a.status_code, 409)
        self.assertEqual(TransactionValorix.objects.count(), 1)
